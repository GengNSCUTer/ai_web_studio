"""受用户和工作区权限保护的 Markdown/DOCX 文档草稿交付。"""

from __future__ import annotations

import hashlib
import json
import os
from pathlib import Path
from tempfile import NamedTemporaryFile
from uuid import uuid4

from docx import Document
from docx.shared import Pt
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.project_file import ProjectFile
from app.repositories.project_file_repo import ProjectFileRepository
from app.repositories.project_repo import ProjectRepository
from app.models.tool_config import WorkspaceAgentPolicy
from app.schemas.document_delivery import DocumentDraftCreate, DocumentDraftResponse


class DocumentDeliveryError(ValueError):
    pass


class DocumentDeliveryService:
    """把聊天生成的文档先保存为独立草稿，不覆盖已有文件。"""

    MAX_REFERENCES = 100

    def __init__(self, db: Session) -> None:
        self.db = db
        self.file_repo = ProjectFileRepository(db)
        self.project_repo = ProjectRepository(db)

    @staticmethod
    def _safe_file_name(file_name: str, file_format: str) -> str:
        name = Path(file_name).name.strip()
        if not name or name in {".", ".."}:
            raise DocumentDeliveryError("文件名不能为空。")
        expected_suffix = ".md" if file_format == "markdown" else ".docx"
        if Path(name).suffix.lower() != expected_suffix:
            name = f"{Path(name).stem}{expected_suffix}"
        if len(name) > 180:
            raise DocumentDeliveryError("文件名过长。")
        return name

    @staticmethod
    def _markdown_with_sources(content: str, references: list[str]) -> str:
        if not references:
            return content.rstrip() + "\n"
        sources = "\n".join(f"- {item}" for item in references)
        return f"{content.rstrip()}\n\n## 来源\n\n{sources}\n"

    @staticmethod
    def _render_docx(content: str, references: list[str], target: Path) -> None:
        document = Document()
        normal = document.styles["Normal"]
        normal.font.name = "Microsoft YaHei"
        normal.font.size = Pt(10.5)
        for paragraph in content.splitlines():
            text = paragraph.strip()
            if not text:
                document.add_paragraph("")
            elif text.startswith("# "):
                document.add_heading(text[2:].strip(), level=1)
            elif text.startswith("## "):
                document.add_heading(text[3:].strip(), level=2)
            elif text.startswith("### "):
                document.add_heading(text[4:].strip(), level=3)
            elif text.startswith("- "):
                document.add_paragraph(text[2:].strip(), style="List Bullet")
            else:
                document.add_paragraph(text)
        if references:
            document.add_heading("来源", level=2)
            for reference in references:
                document.add_paragraph(reference, style="List Bullet")
        document.save(str(target))

    def create_draft(
        self,
        *,
        user_id: str,
        project_id: str,
        payload: DocumentDraftCreate,
    ) -> DocumentDraftResponse:
        if not self.project_repo.get_by_user(project_id, user_id):
            raise DocumentDeliveryError("工作区不存在或无权访问。")
        policy = self.db.get(WorkspaceAgentPolicy, project_id)
        if policy and policy.permission_mode == "read_only":
            raise DocumentDeliveryError("工作区处于只读模式，不能创建草稿；请先调整工作区权限。")
        references = list(dict.fromkeys(item.strip() for item in payload.source_references if item.strip()))[: self.MAX_REFERENCES]
        file_name = self._safe_file_name(payload.file_name, payload.format)
        content = payload.content.strip()
        if not content:
            raise DocumentDeliveryError("文档内容不能为空。")
        user_root = Path(settings.upload_dir) / user_id / "project_documents" / project_id
        user_root.mkdir(parents=True, exist_ok=True)
        storage_name = f"{uuid4()}{Path(file_name).suffix.lower()}"
        target = user_root / storage_name
        rendered_markdown = self._markdown_with_sources(content, references)
        try:
            with NamedTemporaryFile(dir=user_root, prefix=".draft.", suffix=".tmp", delete=False) as temp:
                temp_path = Path(temp.name)
            if payload.format == "markdown":
                temp_path.write_text(rendered_markdown, encoding="utf-8")
                parsed_text = rendered_markdown
                mime_type = "text/markdown"
            else:
                self._render_docx(content, references, temp_path)
                parsed_text = content
                mime_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            os.replace(temp_path, target)
        except Exception:
            if "temp_path" in locals():
                temp_path.unlink(missing_ok=True)
            raise

        storage_key = f"{user_id}/project_documents/{project_id}/{storage_name}"
        try:
            item = ProjectFile(
                project_id=project_id,
                user_id=user_id,
                kind="generated_document",
                file_name=file_name,
                mime_type=mime_type,
                file_size=target.stat().st_size,
                storage_key=storage_key,
                parsed_text=parsed_text,
            )
            item = self.file_repo.save(item)
        except Exception:
            target.unlink(missing_ok=True)
            raise
        digest = hashlib.sha256(target.read_bytes()).hexdigest()
        return DocumentDraftResponse(
            file_id=item.id,
            project_id=project_id,
            file_name=item.file_name,
            format=payload.format,
            file_size=item.file_size or 0,
            content_hash=digest,
            source_references=references,
            download_url=f"/api/projects/{project_id}/files/{item.id}/download",
            created_at=item.created_at,
        )

    def resolve_download_path(self, *, user_id: str, project_id: str, file_id: str) -> tuple[ProjectFile, Path] | None:
        if not self.project_repo.get_by_user(project_id, user_id):
            return None
        item = self.file_repo.get_by_user(file_id, user_id)
        if not item or item.project_id != project_id:
            return None
        root = (Path(settings.upload_dir) / user_id).resolve()
        path = (Path(settings.upload_dir) / item.storage_key).resolve()
        if root not in path.parents or not path.is_file():
            return None
        return item, path
